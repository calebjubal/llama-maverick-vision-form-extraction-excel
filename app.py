from dotenv import load_dotenv

load_dotenv()

import streamlit as st
import pandas as pd
import json
import base64
import hashlib
from groq import Groq
from openpyxl import load_workbook
from PIL import Image
import io

# ================= CONFIG ================= #
MAX_IMAGES = 20
EDITOR_KEY = "review_editor_table"
EDITOR_SIGNATURE_KEY = "review_editor_signature"
EDITOR_DF_KEY = "review_editor_dataframe"

TARGET_FIELDS = [
    "Workshop Code",
    "Workshop Name",
    "Location",
    "Claim No",
    "Dealer Lot",
    "Mileage",
    "Repair Date",
    "Registration Date",
    "Casual Part No",
    "Casual Part Name",
    "VIN Last 6",
    "Model Code",
    "Dealer Observation"
]


def _normalize_label(value):
    if value is None:
        return ""
    return "".join(ch for ch in str(value).lower() if ch.isalnum())


def detect_sheet_headers(workbook_bytes, sheet_name):
    headers = []
    buffer = io.BytesIO(workbook_bytes)
    wb = None
    try:
        wb = load_workbook(buffer, read_only=True, data_only=True)
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows(values_only=True):
                row_values = [
                    str(cell).strip() if cell is not None else ""
                    for cell in row
                ]
                if any(row_values):
                    headers = row_values
                    break
    finally:
        if wb:
            wb.close()
        buffer.close()

    while headers and headers[-1] == "":
        headers.pop()

    return headers


def build_header_lookup(headers):
    lookup = {}
    for idx, header in enumerate(headers):
        normalized = _normalize_label(header)
        if normalized and normalized not in lookup:
            lookup[normalized] = idx
    return lookup


def build_row_for_headers(row_data, headers, header_lookup):
    if header_lookup:
        row_length = len(headers)
        row_values = ["" for _ in range(row_length)]
        missing_fields = []

        for field in TARGET_FIELDS:
            normalized_field = _normalize_label(field)
            value = row_data.get(field, "")
            if normalized_field in header_lookup:
                row_values[header_lookup[normalized_field]] = value
            else:
                missing_fields.append(field)

        if missing_fields:
            for field in missing_fields:
                row_values.append(row_data.get(field, ""))

        return row_values, missing_fields

    fallback_row = [row_data.get(field, "") for field in TARGET_FIELDS]
    return fallback_row, []


def build_records_signature(records):
    serialized = json.dumps(records, sort_keys=True)
    return hashlib.sha256(serialized.encode()).hexdigest()

# ================= INIT ================= #
client = Groq(api_key=st.secrets["GROQ_API_KEY"])

# ================= LOCAL TESTING ================= #
# client = Groq()

# ================= PROMPT ================= #
def build_prompt():
    return f"""
You are a vision-based information extraction engine.

Extract information from the image and return ONLY valid JSON.
Do NOT include explanations.
Do NOT include markdown.
Do NOT hallucinate values.
If a field is missing or unclear, return an empty string.

TARGET SCHEMA (keys must match EXACTLY):

{json.dumps({k: "" for k in TARGET_FIELDS}, indent=2)}

RULES:
- extract direct place/city name for the location (e.g., "PUNALUR-SRV, PUNALUR" -> "PUNALUR")
- Dates must be dd-mm-yy
- Mileage must be an integer
- VIN Last 6 must be numeric
- Dealer Observation must be UPPERCASE
"""

# ================= IMAGE → STRUCTURED ================= #
def image_to_structured(image_bytes):
    image_b64 = base64.b64encode(image_bytes).decode()

    completion = client.chat.completions.create(
        model="meta-llama/llama-4-maverick-17b-128e-instruct",
        messages=[
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": build_prompt()},
                    {
                        "type": "image_url",
                        "image_url": {
                            "url": f"data:image/png;base64,{image_b64}"
                        }
                    }
                ]
            }
        ],
        temperature=0,
        max_completion_tokens=1024,
        stream=True
    )

    full_response = ""
    for chunk in completion:
        delta = chunk.choices[0].delta
        if delta and delta.content:
            full_response += delta.content

    return json.loads(full_response)

# ================= STREAMLIT UI ================= #
st.set_page_config(
    page_title="Vision Form Extraction (Groq Maverick)",
    layout="wide"
)

st.title("📸 Image → Structured Data → Excel (Groq Maverick)")

images = st.file_uploader(
    "Upload form images (max 20)",
    type=["png", "jpg", "jpeg"],
    accept_multiple_files=True
)

excel = st.file_uploader(
    "Upload Excel file (.xlsx)",
    type=["xlsx"]
)

if images and len(images) > MAX_IMAGES:
    st.error("Maximum 20 images allowed")
    st.stop()

if not images or not excel:
    st.info("Upload images and an Excel file to continue")
    st.stop()

if "uploaded_excel_name" not in st.session_state:
    st.session_state["uploaded_excel_name"] = None
    st.session_state["excel_bytes_cache"] = None
    st.session_state["updated_workbook"] = None
    st.session_state["updated_sheet"] = None
    st.session_state["update_message"] = None

if st.session_state["uploaded_excel_name"] != excel.name:
    st.session_state["uploaded_excel_name"] = excel.name
    st.session_state["excel_bytes_cache"] = excel.getvalue()
    st.session_state["updated_workbook"] = None
    st.session_state["updated_sheet"] = None
    st.session_state["update_message"] = None
elif st.session_state["excel_bytes_cache"] is None:
    st.session_state["excel_bytes_cache"] = excel.getvalue()

excel_bytes = st.session_state["excel_bytes_cache"]

preview_buffer = io.BytesIO(excel_bytes)
wb_preview = load_workbook(preview_buffer, read_only=True)
excel_sheetnames = wb_preview.sheetnames
wb_preview.close()
preview_buffer.close()

if not excel_sheetnames:
    st.error("No worksheets found in the uploaded workbook")
    st.stop()

records = []
image_map = {}
failed = []

# ================= PROCESS ================= #
with st.spinner("Processing images with Groq Maverick Vision Model..."):
    for idx, img in enumerate(images):
        try:
            img_bytes = img.read()
            structured = image_to_structured(img_bytes)

            record = structured.copy()
            record["Row ID"] = idx
            record["Source Image"] = img.name

            records.append(record)
            image_map[idx] = img_bytes

        except Exception:
            failed.append(img.name)

# ================= SUMMARY ================= #
st.subheader("📊 Processing Summary")
st.success(f"Extracted records: {len(records)}")
st.error(f"Failed images: {len(failed)}")

if failed:
    with st.expander("❌ Failed Images"):
        for f in failed:
            st.write(f)

# ================= REVIEW TABLE ================= #
if records:
    df = pd.DataFrame(records)
    ordered_cols = ["Row ID", "Source Image"] + TARGET_FIELDS
    df = df[ordered_cols]

    records_signature = build_records_signature(records)

    if st.session_state.get(EDITOR_SIGNATURE_KEY) != records_signature:
        st.session_state[EDITOR_SIGNATURE_KEY] = records_signature
        st.session_state[EDITOR_DF_KEY] = df.copy()
        if EDITOR_KEY in st.session_state:
            del st.session_state[EDITOR_KEY]
    elif EDITOR_DF_KEY not in st.session_state:
        st.session_state[EDITOR_DF_KEY] = df.copy()

    st.subheader("✏️ Review & Edit Extracted Data")

    edited_df = st.data_editor(
        st.session_state[EDITOR_DF_KEY],
        num_rows="dynamic",
        use_container_width=True,
        key=EDITOR_KEY
    )

    st.session_state[EDITOR_DF_KEY] = edited_df.copy()

    # ================= IMAGE PREVIEW ================= #
    st.subheader("🖼️ Preview Source Image")

    selected_row = st.selectbox(
        "Select Row ID to preview image",
        options=edited_df["Row ID"].tolist()
    )

    if selected_row in image_map:
        st.image(
            image_map[selected_row],
            caption=f"Source Image for Row ID {selected_row}",
            use_column_width=True
        )

    st.subheader("📎 Excel Destination")
    selected_sheet = st.selectbox(
        "Select worksheet to append rows",
        options=excel_sheetnames,
        help="All reviewed rows will be appended to this sheet."
    )

    sheet_headers = detect_sheet_headers(excel_bytes, selected_sheet)
    header_lookup = build_header_lookup(sheet_headers)
    missing_header_fields = [
        field for field in TARGET_FIELDS
        if _normalize_label(field) not in header_lookup
    ]

    if sheet_headers:
        pretty_headers = [h if h else "<blank>" for h in sheet_headers]
        st.caption(
            f"Detected headers in '{selected_sheet}': {', '.join(pretty_headers)}"
        )
    else:
        st.warning(
            "Could not detect headers in this worksheet. Data will follow the default field order."
        )

    if missing_header_fields:
        st.info(
            "Columns not found in sheet: " + ", ".join(missing_header_fields)
        )

    # ================= SUBMIT ================= #
    if st.button("✅ Submit to Excel"):
        with st.spinner("Updating workbook..."):
            workbook_buffer = io.BytesIO(excel_bytes)
            wb = load_workbook(workbook_buffer)

            if selected_sheet not in wb.sheetnames:
                st.error("Selected worksheet is not available in the uploaded workbook")
            else:
                ws = wb[selected_sheet]
                cumulative_missing = set()

                for _, row in edited_df.iterrows():
                    row_values, row_missing = build_row_for_headers(
                        row.to_dict(),
                        sheet_headers,
                        header_lookup
                    )
                    cumulative_missing.update(row_missing)
                    ws.append(row_values)

                output_buffer = io.BytesIO()
                wb.save(output_buffer)
                output_buffer.seek(0)
                updated_bytes = output_buffer.getvalue()
                output_buffer.close()

                st.session_state["excel_bytes_cache"] = updated_bytes
                st.session_state["updated_workbook"] = updated_bytes
                st.session_state["updated_sheet"] = selected_sheet
                message = f"Excel sheet '{selected_sheet}' updated successfully."
                if cumulative_missing:
                    message += " Columns added at the end for: " + \
                        ", ".join(sorted(cumulative_missing))
                st.session_state["update_message"] = message

            wb.close()
            workbook_buffer.close()

    if st.session_state.get("update_message"):
        st.success(st.session_state["update_message"])

    if st.session_state.get("updated_workbook"):
        st.download_button(
            "⬇️ Download updated workbook",
            data=st.session_state["updated_workbook"],
            file_name=f"updated_{excel.name}",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

else:
    st.warning("No valid data extracted")
